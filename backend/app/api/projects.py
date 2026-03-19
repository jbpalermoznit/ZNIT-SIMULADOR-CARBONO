from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
from app.core.database import get_db
from app.core.auth import get_current_user
from app.models.user import User
from app.models.project import Project
from app.models.abc_item import AbcCurve, AbcItem
from app.schemas.project import ProjectCreate, ProjectResponse, ProjectWithStats
from app.schemas.abc_item import AbcCurveResponse, AbcItemResponse, UploadResult
from app.services.parser import parse_abc_file

router = APIRouter(prefix="/projects", tags=["projects"])


# ---------------------------------------------------------------------------
# Projetos
# ---------------------------------------------------------------------------

@router.get("", response_model=list[ProjectWithStats])
def list_projects(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from app.models.scenario import Scenario, ScenarioResult

    projects = (
        db.query(Project)
        .filter(Project.company_id == current_user.company_id)
        .order_by(Project.created_at.desc())
        .all()
    )

    results = []
    for p in projects:
        # Count scenarios
        scenarios_count = db.query(Scenario).filter(Scenario.project_id == p.id).count()

        # Count items
        curve = (
            db.query(AbcCurve)
            .filter(AbcCurve.project_id == p.id)
            .order_by(AbcCurve.imported_at.desc())
            .first()
        )
        items_count = (
            db.query(AbcItem).filter(AbcItem.abc_curve_id == curve.id).count()
            if curve else 0
        )

        # Get base scenario result
        base = (
            db.query(Scenario)
            .filter(Scenario.project_id == p.id, Scenario.is_base == True)
            .first()
        )
        result = None
        if base:
            result = (
                db.query(ScenarioResult)
                .filter(ScenarioResult.scenario_id == base.id)
                .first()
            )

        results.append(ProjectWithStats(
            id=p.id,
            company_id=p.company_id,
            name=p.name,
            client_name=p.client_name,
            address=p.address,
            total_area_m2=p.total_area_m2,
            building_type=p.building_type,
            status=p.status,
            created_at=p.created_at,
            total_tco2e=result.total_tco2e if result else None,
            intensity_kgco2e_per_m2=(
                (result.total_tco2e * 1000 / p.total_area_m2)
                if result and p.total_area_m2
                else None
            ),
            coverage_pct=result.coverage_pct if result else None,
            scenarios_count=scenarios_count,
            items_count=items_count,
        ))
    return results


@router.post("", response_model=ProjectResponse, status_code=201)
def create_project(
    body: ProjectCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = Project(company_id=current_user.company_id, **body.model_dump())
    db.add(project)
    db.commit()
    db.refresh(project)
    return project


@router.get("/{project_id}", response_model=ProjectResponse)
def get_project(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = _get_project_or_404(project_id, current_user.company_id, db)
    return project


# ---------------------------------------------------------------------------
# Upload Curva ABC
# ---------------------------------------------------------------------------

@router.post("/{project_id}/upload-abc", response_model=UploadResult)
async def upload_abc(
    project_id: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    project = _get_project_or_404(project_id, current_user.company_id, db)

    # Validar extensão
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Formato inválido. Use .xlsx ou .xlsm")

    content = await file.read()

    # Executar parser
    try:
        result = parse_abc_file(content, file.filename)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    # Salvar AbcCurve
    curve = AbcCurve(
        project_id=project.id,
        file_name=file.filename,
        imported_by_user_id=current_user.id,
        total_items=len(result.items),
        total_cost=result.total_cost,
    )
    db.add(curve)
    db.flush()  # obter curve.id antes de criar os itens

    # Salvar itens
    db_items = [
        AbcItem(
            abc_curve_id=curve.id,
            cost_code=item.cost_code,
            description=item.description,
            adf=item.adf,
            quantity=item.quantity,
            unit=item.unit,
            unit_cost=item.unit_cost,
            total_cost=item.total_cost,
            supplier=item.supplier,
            cost_pct=item.cost_pct,
            cumulative_pct=item.cumulative_pct,
            abc_class=item.abc_class,
            item_type=item.item_type,
            item_order=item.order,
            mapping_status=item.mapping_status,
            classification_note=item.classification_note,
        )
        for item in result.items
    ]
    db.bulk_save_objects(db_items)
    db.commit()

    # Resumo por tipo e classe
    type_summary = {}
    class_summary = {}
    for item in result.items:
        type_summary[item.item_type] = type_summary.get(item.item_type, 0) + 1
        class_summary[item.abc_class] = class_summary.get(item.abc_class, 0) + 1

    return UploadResult(
        abc_curve_id=curve.id,
        file_name=file.filename,
        total_items=len(result.items),
        total_cost=result.total_cost,
        type_summary=type_summary,
        class_summary=class_summary,
        warnings=result.warnings,
    )


# ---------------------------------------------------------------------------
# Itens da Curva ABC
# ---------------------------------------------------------------------------

@router.get("/{project_id}/abc-items", response_model=list[AbcItemResponse])
def list_abc_items(
    project_id: str,
    item_type: Optional[str] = None,
    abc_class: Optional[str] = None,
    mapping_status: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    from app.models.item_mapping import ItemMapping

    _get_project_or_404(project_id, current_user.company_id, db)

    # Buscar a curva mais recente do projeto
    curve = (
        db.query(AbcCurve)
        .filter(AbcCurve.project_id == project_id)
        .order_by(AbcCurve.imported_at.desc())
        .first()
    )
    if not curve:
        return []

    query = db.query(AbcItem).filter(AbcItem.abc_curve_id == curve.id)

    if item_type:
        query = query.filter(AbcItem.item_type == item_type)
    if abc_class:
        query = query.filter(AbcItem.abc_class == abc_class)
    if mapping_status:
        query = query.filter(AbcItem.mapping_status == mapping_status)

    items = query.order_by(AbcItem.item_order).all()

    # Enriquecer com dados do mapping
    item_ids = [i.id for i in items]
    mappings = (
        db.query(ItemMapping)
        .filter(ItemMapping.abc_item_id.in_(item_ids))
        .all()
    ) if item_ids else []
    mapping_by_item = {m.abc_item_id: m for m in mappings}

    results = []
    for item in items:
        m = mapping_by_item.get(item.id)
        results.append(AbcItemResponse(
            id=item.id,
            abc_curve_id=item.abc_curve_id,
            cost_code=item.cost_code,
            description=item.description,
            adf=item.adf,
            quantity=item.quantity,
            unit=item.unit,
            unit_cost=item.unit_cost,
            total_cost=item.total_cost,
            supplier=item.supplier,
            cost_pct=item.cost_pct,
            cumulative_pct=item.cumulative_pct,
            abc_class=item.abc_class,
            item_type=item.item_type,
            item_order=item.item_order,
            mapping_status=item.mapping_status,
            classification_note=item.classification_note,
            parent_item_id=item.parent_item_id,
            factor_name=m.factor_name if m else None,
            factor_value=m.factor_value if m else None,
            factor_unit=m.factor_unit if m else None,
            source_tier=m.source_tier if m else None,
            confidence=m.confidence if m else None,
        ))
    return results


@router.get("/{project_id}/abc-curves", response_model=list[AbcCurveResponse])
def list_abc_curves(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _get_project_or_404(project_id, current_user.company_id, db)
    return (
        db.query(AbcCurve)
        .filter(AbcCurve.project_id == project_id)
        .order_by(AbcCurve.imported_at.desc())
        .all()
    )


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

TYPE_LABELS = {
    "A": "Material Direto",
    "B": "Mão de Obra",
    "C": "Item Agrupado",
    "D": "Material Embutido",
    "E": "Equipamento/Locação",
    "F": "Administrativo/Indireto",
}

TYPE_PARAMETRIZATION = {
    "A": "Buscar fator de emissão no banco (CECarbon, Ecoinvent ou GHG Protocol). Unidade do fator deve ser compatível com a unidade do item.",
    "B": "Excluir do inventário — mão de obra sem emissão direta. Justificativa: escopo operacional GHG Protocol.",
    "C": "Decompor em sub-itens (material, equipamento, mão de obra) com % de custo. Cada sub-item recebe seu próprio fator.",
    "D": "Verificar dupla contagem com itens Tipo A. Se o material já está contabilizado em outro item, excluir este.",
    "E": "Configurar: tipo combustível (diesel/elétrico), consumo por hora (L/h ou kWh/h), fator de emissão do combustível.",
    "F": "Excluir do inventário — custo administrativo sem emissão direta. Justificativa: fora do escopo operacional.",
}

STATUS_LABELS = {
    "auto": "Mapeado (automático)",
    "manual": "Sugerido (revisar)",
    "pending": "Pendente",
    "blocked": "Bloqueado",
    "excluded": "Excluído",
}


@router.get("/{project_id}/export-items")
def export_items_excel(
    project_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Export all items to Excel with factor suggestions and parametrization guidance."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from app.models.item_mapping import ItemMapping

    _get_project_or_404(project_id, current_user.company_id, db)

    project = db.query(Project).filter(Project.id == project_id).first()

    curve = (
        db.query(AbcCurve)
        .filter(AbcCurve.project_id == project_id)
        .order_by(AbcCurve.imported_at.desc())
        .first()
    )
    if not curve:
        raise HTTPException(404, "Nenhuma curva ABC importada")

    items = (
        db.query(AbcItem)
        .filter(AbcItem.abc_curve_id == curve.id)
        .order_by(AbcItem.item_order)
        .all()
    )

    # Load mappings
    item_ids = [i.id for i in items]
    mappings = (
        db.query(ItemMapping)
        .filter(ItemMapping.abc_item_id.in_(item_ids))
        .all()
    ) if item_ids else []
    mapping_by_item = {m.abc_item_id: m for m in mappings}

    # Create workbook
    wb = openpyxl.Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="1d7a6b", end_color="1d7a6b", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E0E4E3"),
        right=Side(style="thin", color="E0E4E3"),
        top=Side(style="thin", color="E0E4E3"),
        bottom=Side(style="thin", color="E0E4E3"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    status_fills = {
        "auto": PatternFill(start_color="E6F3EE", end_color="E6F3EE", fill_type="solid"),
        "manual": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        "pending": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "excluded": PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid"),
    }
    type_tab_colors = {
        "A": "1d7a6b", "B": "92400e", "C": "7c3aed",
        "D": "be123c", "E": "1e40af", "F": "4b5563",
    }

    headers = [
        "Código", "Descrição", "Quantidade", "Unidade", "Custo Unitário",
        "Custo Total", "% Custo", "% Acumulado", "Classe Pareto",
        "Tipo", "Tipo (Descrição)", "Status",
        "Fator Sugerido", "Valor Fator", "Unidade Fator", "Fonte",
        "Confiança", "Score Match",
        "Sugestão de Parametrização",
        "Ação do Usuário (preencher)", "Fator Manual (preencher)", "Justificativa (preencher)",
    ]
    widths = [15, 45, 12, 8, 12, 15, 8, 8, 8, 6, 20, 18, 30, 12, 15, 15, 12, 8, 45, 20, 20, 30]

    def _write_sheet(ws, sheet_items):
        """Write headers and data rows to a worksheet."""
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row_idx, item in enumerate(sheet_items, 2):
            m = mapping_by_item.get(item.id)
            is_sub = bool(item.parent_item_id)
            data = [
                ("→ " if is_sub else "") + (item.cost_code or ""),
                ("  ↳ " if is_sub else "") + (item.description or ""),
                item.quantity, item.unit, item.unit_cost, item.total_cost,
                round(item.cost_pct * 100, 2) if item.cost_pct else 0,
                round(item.cumulative_pct * 100, 2) if item.cumulative_pct else 0,
                item.abc_class or "", item.item_type,
                TYPE_LABELS.get(item.item_type, ""),
                STATUS_LABELS.get(item.mapping_status, item.mapping_status),
                m.factor_name if m else "", m.factor_value if m else "",
                m.factor_unit if m else "", m.source_tier if m else "",
                m.confidence if m else "",
                round(m.similarity_score * 100, 0) if m and m.similarity_score else "",
                TYPE_PARAMETRIZATION.get(item.item_type, ""),
                "", "", "",
            ]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = thin_border
                cell.alignment = wrap
                if col == 12:
                    fill = status_fills.get(item.mapping_status)
                    if fill:
                        cell.fill = fill

        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        if sheet_items:
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(sheet_items) + 1}"

    # --- Sheet 1: Resumo (all items) ---
    ws_all = wb.active
    ws_all.title = "Todos os Itens"
    _write_sheet(ws_all, items)

    # --- Sheets per type ---
    for type_code in ["A", "B", "C", "D", "E", "F"]:
        type_items = [i for i in items if i.item_type == type_code]
        # Include sub-items of this type's parent items
        parent_ids = {i.id for i in type_items}
        sub_items = [i for i in items if i.parent_item_id and i.parent_item_id in parent_ids]
        combined = type_items + [s for s in sub_items if s not in type_items]
        if not combined:
            continue
        label_clean = TYPE_LABELS[type_code].replace("/", "-")
        sheet_name = f"Tipo {type_code} - {label_clean}"
        if len(sheet_name) > 31:
            sheet_name = sheet_name[:31]
        ws_type = wb.create_sheet(sheet_name)
        ws_type.sheet_properties.tabColor = type_tab_colors.get(type_code, "808080")
        _write_sheet(ws_type, combined)

    # --- Legend sheet ---
    ws_leg = wb.create_sheet("Legenda")
    ws_leg.cell(row=1, column=1, value="Tipos de Item").font = Font(bold=True, size=12)
    for i, (tipo, desc) in enumerate(TYPE_LABELS.items(), 3):
        ws_leg.cell(row=i, column=1, value=f"Tipo {tipo}").font = Font(bold=True)
        ws_leg.cell(row=i, column=2, value=desc)
        ws_leg.cell(row=i, column=3, value=TYPE_PARAMETRIZATION[tipo])
    ws_leg.column_dimensions["A"].width = 10
    ws_leg.column_dimensions["B"].width = 25
    ws_leg.column_dimensions["C"].width = 80

    ws_leg.cell(row=11, column=1, value="Status").font = Font(bold=True, size=12)
    for i, (status, label) in enumerate(STATUS_LABELS.items(), 13):
        ws_leg.cell(row=i, column=1, value=status)
        ws_leg.cell(row=i, column=2, value=label)

    ws_leg.cell(row=20, column=1, value="Fontes de Fatores").font = Font(bold=True, size=12)
    sources = [
        ("ghg_protocol", "GHG Protocol Brasil — combustíveis e energia"),
        ("cecarbon", "CECarbon — materiais de construção brasileiros (120+ materiais)"),
        ("ecoinvent", "Ecoinvent — base global de LCA (fatores europeus/globais)"),
        ("user_custom", "Fator manual inserido pelo analista"),
        ("excluded", "Item excluído do inventário"),
    ]
    for i, (src, desc) in enumerate(sources, 22):
        ws_leg.cell(row=i, column=1, value=src)
        ws_leg.cell(row=i, column=2, value=desc)

    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    from datetime import datetime
    from urllib.parse import quote
    date_str = datetime.now().strftime("%Y-%m-%d")
    project_name = project.name.replace(" ", "_").replace("/", "-") if project else "Projeto"
    filename = f"ZNIT_{project_name}_{date_str}.xlsx"
    filename_encoded = quote(filename)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{filename_encoded}",
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


# ---------------------------------------------------------------------------
# Power BI — JSON endpoints (flat tabular data)
# Accepts JWT auth OR api_key query param for Power BI access
# ---------------------------------------------------------------------------

from fastapi import Query, Request

def _get_powerbi_auth(
    request: Request,
    api_key: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Authenticate via JWT or API key for Power BI endpoints."""
    from app.core.config import settings

    # Try API key first
    if api_key and api_key == settings.POWERBI_API_KEY:
        return True

    # Try JWT
    auth_header = request.headers.get("authorization", "")
    if auth_header.startswith("Bearer "):
        try:
            from app.core.auth import decode_token
            token = auth_header.split(" ")[1]
            payload = decode_token(token)
            if payload:
                return True
        except Exception:
            pass

    raise HTTPException(401, "API key inválida ou token expirado")


@router.get("/{project_id}/powerbi/items")
def powerbi_items(
    project_id: str,
    db: Session = Depends(get_db),
    _auth: bool = Depends(_get_powerbi_auth),
):
    """Flat JSON table of all items with emissions — for Power BI / Data Studio."""
    from app.models.item_mapping import ItemMapping
    from app.services.calculator import normalize_unit, get_conversion_factor

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(404, "Projeto não encontrado")

    curve = (
        db.query(AbcCurve)
        .filter(AbcCurve.project_id == project_id)
        .order_by(AbcCurve.imported_at.desc())
        .first()
    )
    if not curve:
        return []

    items = db.query(AbcItem).filter(AbcItem.abc_curve_id == curve.id).order_by(AbcItem.item_order).all()
    item_ids = [i.id for i in items]
    mappings = db.query(ItemMapping).filter(ItemMapping.abc_item_id.in_(item_ids)).all() if item_ids else []
    mapping_by_item = {m.abc_item_id: m for m in mappings}

    rows = []
    for item in items:
        m = mapping_by_item.get(item.id)
        conv = get_conversion_factor(item.unit, m.factor_unit) if m and m.factor_unit else 1
        emission_kgco2e = (item.quantity or 0) * (m.factor_value or 0) * conv if m and m.factor_value and conv > 0 else 0

        rows.append({
            "projeto": project.name if project else "",
            "area_m2": project.total_area_m2 if project else None,
            "codigo": item.cost_code,
            "descricao": item.description,
            "quantidade": item.quantity,
            "unidade": item.unit,
            "custo_unitario": item.unit_cost,
            "custo_total": item.total_cost,
            "custo_pct": round(item.cost_pct * 100, 4) if item.cost_pct else 0,
            "custo_acumulado_pct": round(item.cumulative_pct * 100, 4) if item.cumulative_pct else 0,
            "classe_pareto": item.abc_class,
            "tipo": item.item_type,
            "tipo_descricao": TYPE_LABELS.get(item.item_type, ""),
            "status": item.mapping_status,
            "status_descricao": STATUS_LABELS.get(item.mapping_status, item.mapping_status),
            "fator_nome": m.factor_name if m else None,
            "fator_valor": m.factor_value if m else None,
            "fator_unidade": m.factor_unit if m else None,
            "fator_fonte": m.source_tier if m else None,
            "confianca": m.confidence if m else None,
            "score_match": round(m.similarity_score * 100, 1) if m and m.similarity_score else None,
            "emissao_kgco2e": round(emission_kgco2e, 2),
            "emissao_tco2e": round(emission_kgco2e / 1000, 4),
            "intensidade_kgco2e_m2": round(emission_kgco2e / project.total_area_m2, 4) if project and project.total_area_m2 else None,
            "eh_sub_item": bool(item.parent_item_id),
            "item_pai_id": item.parent_item_id,
        })
    return rows


@router.get("/{project_id}/powerbi/scenarios")
def powerbi_scenarios(
    project_id: str,
    db: Session = Depends(get_db),
    _auth: bool = Depends(_get_powerbi_auth),
):
    """Flat JSON table of scenarios with results — for Power BI."""
    from app.models.scenario import Scenario, ScenarioResult

    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise HTTPException(404, "Projeto não encontrado")

    scenarios = db.query(Scenario).filter(Scenario.project_id == project_id).all()

    rows = []
    for scen in scenarios:
        result = db.query(ScenarioResult).filter(ScenarioResult.scenario_id == scen.id).first()
        rows.append({
            "projeto": project.name if project else "",
            "cenario": scen.name,
            "cenario_id": scen.id,
            "eh_base": scen.is_base,
            "status": scen.status,
            "total_kgco2e": result.total_kgco2e if result else 0,
            "total_tco2e": result.total_tco2e if result else 0,
            "intensidade_tco2e_m2": result.intensity_per_m2 if result else 0,
            "scope1_kgco2e": result.scope1_kgco2e if result else 0,
            "scope2_kgco2e": result.scope2_kgco2e if result else 0,
            "scope3_materiais_kgco2e": result.scope3_materials_kgco2e if result else 0,
            "scope3_logistica_kgco2e": result.scope3_logistics_kgco2e if result else 0,
            "itens_total": result.items_total if result else 0,
            "itens_mapeados": result.items_mapped if result else 0,
            "itens_excluidos": result.items_excluded if result else 0,
            "cobertura_pct": result.coverage_pct if result else 0,
            "calculado_em": str(result.calculated_at) if result and result.calculated_at else None,
        })
    return rows


@router.get("/{project_id}/powerbi/emissoes-por-categoria")
def powerbi_emissions_by_category(
    project_id: str,
    db: Session = Depends(get_db),
    _auth: bool = Depends(_get_powerbi_auth),
):
    """Emissions grouped by material category — for Power BI charts."""
    from app.models.item_mapping import ItemMapping
    from app.models.scenario import Scenario, ScenarioItem

    # Get base scenario
    base = db.query(Scenario).filter(
        Scenario.project_id == project_id, Scenario.is_base == True
    ).first()
    if not base:
        return []

    items = db.query(ScenarioItem).filter(ScenarioItem.scenario_id == base.id).all()
    curve = db.query(AbcCurve).filter(AbcCurve.project_id == project_id).order_by(AbcCurve.imported_at.desc()).first()
    if not curve:
        return []

    abc_items = {i.id: i for i in db.query(AbcItem).filter(AbcItem.abc_curve_id == curve.id).all()}

    # Group by factor_name (material category)
    categories: dict[str, dict] = {}
    for si in items:
        abc = abc_items.get(si.abc_item_id)
        if not abc or not si.emission_kgco2e or si.emission_kgco2e <= 0:
            continue

        cat_name = si.factor_name or abc.description
        if cat_name not in categories:
            categories[cat_name] = {
                "categoria": cat_name,
                "fonte": si.source_tier or "",
                "tipo_item": abc.item_type,
                "itens_count": 0,
                "emissao_kgco2e": 0,
                "emissao_tco2e": 0,
                "custo_total": 0,
            }
        categories[cat_name]["itens_count"] += 1
        categories[cat_name]["emissao_kgco2e"] += si.emission_kgco2e
        categories[cat_name]["emissao_tco2e"] += si.emission_kgco2e / 1000
        categories[cat_name]["custo_total"] += abc.total_cost or 0

    rows = sorted(categories.values(), key=lambda r: r["emissao_kgco2e"], reverse=True)

    # Add percentage
    total = sum(r["emissao_kgco2e"] for r in rows) or 1
    cumulative = 0
    for r in rows:
        r["pct_total"] = round(r["emissao_kgco2e"] / total * 100, 2)
        cumulative += r["pct_total"]
        r["pct_acumulado"] = round(cumulative, 2)
        r["emissao_kgco2e"] = round(r["emissao_kgco2e"], 2)
        r["emissao_tco2e"] = round(r["emissao_tco2e"], 4)
        r["custo_total"] = round(r["custo_total"], 2)

    return rows


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _get_project_or_404(project_id: str, company_id: str, db: Session) -> Project:
    project = (
        db.query(Project)
        .filter(Project.id == project_id, Project.company_id == company_id)
        .first()
    )
    if not project:
        raise HTTPException(status_code=404, detail="Projeto não encontrado")
    return project
